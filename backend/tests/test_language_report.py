import io
import json
import zipfile

from fastapi.testclient import TestClient
from openpyxl import load_workbook
from PIL import Image

from app import config, language, main

GKV = config.DATA_DIR / "examples" / "gkv" / "gkv_taxonomy_l1_l2.json"


def test_language_detection():
    assert language.detect(["Beiträge & Finanzen. Fragen rund um die Höhe der Beiträge"]) == "de"
    assert language.detect(["Billing & Payments. Invoices, charges and payments."]) == "en"
    assert language.detect([]) == "en"
    assert "German" in language.rule("de") and "never in English" in language.rule("de")
    assert language.tidy_label("Befreiung Von Zuzahlungen", "de") == "Befreiung von Zuzahlungen"
    assert language.tidy_label("Krankengeld Nicht Ausgezahlt", "de") == "Krankengeld nicht Ausgezahlt"
    assert language.tidy_label("Von Anfang An", "de") == "Von Anfang an"
    assert language.tidy_label("Plan Change In App", "en") == "Plan Change In App"


def test_new_project_needs_taxonomy_and_detects_language(tmp_db, fake):
    with TestClient(main.app) as client:
        pid = client.post("/api/projects", json={"name": "GKV"}).json()["id"]
        client.post(f"/api/projects/{pid}/activate")
        client.post("/api/import/sample")
        res = client.post("/api/analysis/start")
        assert res.status_code == 409 and "taxonomy" in res.json()["detail"]

        up = client.post("/api/taxonomy/import", files={"file": ("tax.json", GKV.read_bytes(), "application/json")})
        assert up.json()["language"] == "de"
        assert client.get("/api/settings").json()["taxonomy_language"] == "de"
        assert client.get("/api/summary").json()["taxonomy_language"] == "de"

        assert client.put("/api/settings", json={"taxonomy_language": "xx"}).status_code == 400
        assert client.put("/api/settings", json={"taxonomy_language": "en"}).json()["taxonomy_language"] == "en"


def test_prompts_carry_analysis_language(tmp_db, fake):
    with TestClient(main.app) as client:
        seed = json.loads(GKV.read_text())
        client.post("/api/import/sample")
        client.post("/api/taxonomy/import", files={"file": ("tax.json", json.dumps(seed).encode(), "application/json")})
        client.post("/api/analysis/start")
        main.worker._thread.join(timeout=60)
    systems = [p["messages"][0]["content"] for p in fake.payloads if "messages" in p]
    assert systems and all("in German, never in English" in s for s in systems if "taxonomy" in s or "analyst" in s)


def test_localized_exports_and_report(tmp_db, fake, tmp_path):
    with TestClient(main.app) as client:
        client.post("/api/import/sample")
        client.post("/api/analysis/start")
        main.worker._thread.join(timeout=60)

        wb = load_workbook(io.BytesIO(client.get("/api/export/xlsx?lang=de").content))
        assert wb.sheetnames == ["Taxonomie", "Interaktionszuordnungen", "Prüfliste", "Taxonomie-Änderungen"]
        assert wb["Interaktionszuordnungen"]["A1"].value == "Gesprächs-ID"
        assert load_workbook(io.BytesIO(client.get("/api/export/xlsx").content)).sheetnames[0] == "Taxonomy"

        z = zipfile.ZipFile(io.BytesIO(client.get("/api/export/csv?lang=de").content))
        assert "pruefliste.csv" in z.namelist()
        single = client.get("/api/export/csv?sheet=Review%20Queue&lang=de")
        assert single.status_code == 200 and "Grund" in single.content.decode("utf-8-sig").splitlines()[0]

        pdf = client.get("/api/export/report.pdf?lang=de")
        assert pdf.status_code == 200 and pdf.content.startswith(b"%PDF")

        buf = io.BytesIO()
        Image.new("RGB", (400, 100), "red").save(buf, "PNG")
        assert client.post("/api/branding/logo", files={"file": ("logo.png", buf.getvalue(), "image/png")}).json()["logo"]
        assert client.get("/api/branding/logo").headers["content-type"] == "image/png"
        assert client.get("/api/export/report.pdf?lang=en").content.startswith(b"%PDF")
        bad = client.post("/api/branding/logo", files={"file": ("logo.svg", b"<svg/>", "image/svg+xml")})
        assert bad.status_code == 400
        assert client.delete("/api/branding/logo").json()["logo"] is False
