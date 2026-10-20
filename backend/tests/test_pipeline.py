import io
import json
from pathlib import Path

from openpyxl import load_workbook

from app import assignments, db, export, insights, interactions, taxonomy
from app.worker import AnalysisWorker
from conftest import load_sample

APP_DIR = Path(__file__).resolve().parent.parent / "app"


def _run(worker: AnalysisWorker, job="analysis"):
    worker.start(job)
    worker.join(60)
    assert not worker.is_running()


def test_full_analysis_builds_taxonomy(conn, fake):
    db.save_settings(conn, {"batch_size": 10})
    load_sample(conn, 40)
    w = AnalysisWorker()
    _run(w)
    s = insights.summary(conn)
    assert s["processed"] == 40 and s["pending"] == 0
    assert s["l4_nodes"] >= 3 and s["l3_nodes"] >= 1
    run = conn.execute("SELECT * FROM analysis_runs ORDER BY id DESC").fetchone()
    assert run["status"] == "completed"
    kinds = {r[0] for r in conn.execute("SELECT DISTINCT kind FROM events")}
    for k in ("extracting", "intents_extracted", "embedding", "candidate_new", "batch_refinement",
              "cluster_discovered", "category_proposed", "category_approved", "matched_existing"):
        assert k in kinds, k
    # several batches recorded, curve is monotonic
    assert conn.execute("SELECT COUNT(*) FROM batch_stats").fetchone()[0] >= 4
    pts = insights.discovery_curve(conn)["points"]
    assert len(pts) == 40 and all(a["categories"] <= b["categories"] for a, b in zip(pts, pts[1:]))
    # every AI L3/L4 node has a persistent id and a creation record
    for n in taxonomy.nodes(conn, ("active",)):
        if n["level"] >= 3:
            assert taxonomy.list_changes(conn, node_id=n["id"])
    # embeddings are cached, not regenerated
    embeds_before = fake.calls.get("/api/embed", 0)
    _ = w  # second start with nothing pending must be refused
    try:
        w.start()
        raised = False
    except Exception:
        raised = True
    assert raised and fake.calls.get("/api/embed", 0) == embeds_before


def test_multi_intent_rows_and_review_flow(conn, fake):
    db.save_settings(conn, {"batch_size": 25})
    load_sample(conn, 30)
    _run(AnalysisWorker())
    rows = interactions.list_interactions(conn)
    multi = [r for r in rows if len(r["assignments"]) > 1]
    assert multi, "fake model should find multi-intent conversations"
    xrows = export.assignment_rows(conn)
    total_assign = conn.execute("SELECT COUNT(*) FROM assignments").fetchone()[0]
    no_intent = sum(1 for r in rows if not r["assignments"])
    assert len(xrows) == total_assign + no_intent
    # correct one assignment manually; original AI assignment remains in history
    target = conn.execute("SELECT * FROM assignments WHERE l4_id IS NOT NULL LIMIT 1").fetchone()
    other_l2 = "L2-PRD-POR"
    with db.transaction(conn):
        assignments.review(conn, target["id"], "correct", path={"l2_id": other_l2}, notes="Customer meant porting")
    row = conn.execute("SELECT * FROM assignments WHERE id=?", (target["id"],)).fetchone()
    assert row["review_status"] == "corrected" and row["l1_id"] == "L1-PRD" and row["l4_id"] is None
    hist = assignments.history(conn, target["id"])
    assert hist[0]["actor"] == "ai" and hist[-1]["actor"] == "user"
    assert db.loads(row["ai_original_json"])["l4_id"] == target["l4_id"]
    # export reflects the correction
    wb = load_workbook(io.BytesIO(export.to_xlsx(conn)))
    assert wb.sheetnames == ["Taxonomy", "Interaction Assignments", "Review Queue", "Taxonomy Changes"]
    ws = wb["Interaction Assignments"]
    header = [c.value for c in ws[1]]
    notes_col, status_col = header.index("Reviewer Notes"), header.index("Review Status")
    corrected = [r for r in ws.iter_rows(min_row=2, values_only=True) if r[status_col] == "corrected"]
    assert len(corrected) == 1 and "porting" in corrected[0][notes_col]
    assert ws.max_row - 1 == len(xrows)


def test_human_review_survives_restructuring(conn, fake):
    db.save_settings(conn, {"batch_size": 50})
    load_sample(conn, 50)
    _run(AnalysisWorker())
    l4s = [n for n in taxonomy.nodes(conn, ("active",), level=4)]
    by_l2: dict = {}
    for n in l4s:
        by_l2.setdefault(assignments.path_for_node(conn, n["id"])["l2_id"], []).append(n)
    pair = next(v for v in by_l2.values() if len(v) >= 2) if any(len(v) >= 2 for v in by_l2.values()) else None
    if pair is None:  # create a sibling manually so a merge is possible
        n0 = l4s[0]
        with db.transaction(conn):
            sib = taxonomy.manual_create(conn, n0["parent_id"], "Manual Sibling", "test")
        pair = [n0, sib]
    src, tgt = pair[0], pair[1]
    confirmed = conn.execute("SELECT id FROM assignments WHERE l4_id=? LIMIT 1", (src["id"],)).fetchone()
    with db.transaction(conn):
        assignments.review(conn, confirmed["id"], "confirm")
        res = taxonomy.manual_merge(conn, src["id"], tgt["id"], db.load_settings(conn))
    assert taxonomy.get_node(conn, src["id"])["status"] == "merged"
    kept = conn.execute("SELECT * FROM assignments WHERE id=?", (confirmed["id"],)).fetchone()
    assert kept["l4_id"] == src["id"] and kept["review_status"] == "confirmed" and kept["needs_review"] == 1
    assert conn.execute("SELECT COUNT(*) FROM assignments WHERE l4_id=? AND review_status='pending'",
                        (src["id"],)).fetchone()[0] == 0
    assert res["version"] > 1


def test_resume_after_interruption_has_no_duplicates(conn, fake):
    db.save_settings(conn, {"batch_size": 5})
    load_sample(conn, 12)
    w = AnalysisWorker()
    w.start()
    w.pause()
    w.join(30)
    processed = conn.execute("SELECT COUNT(*) FROM interactions WHERE status='processed'").fetchone()[0]
    assert processed < 12
    # simulate a crash in the middle of an interaction
    row = conn.execute("SELECT id FROM interactions WHERE status='pending' ORDER BY seq LIMIT 1").fetchone()
    conn.execute("UPDATE interactions SET status='processing' WHERE id=?", (row["id"],))
    conn.execute("INSERT INTO intents(interaction_id, position, intent_text, created_at) VALUES (?, 0, 'partial', 'x')",
                 (row["id"],))
    db.recover_interrupted(conn)
    assert conn.execute("SELECT COUNT(*) FROM intents WHERE intent_text='partial'").fetchone()[0] == 0
    w.resume()
    w.join(60)
    assert conn.execute("SELECT COUNT(*) FROM interactions WHERE status='processed'").fetchone()[0] == 12
    dup = conn.execute("SELECT interaction_id, position, COUNT(*) c FROM intents GROUP BY interaction_id, position "
                       "HAVING c > 1").fetchall()
    assert not dup
    orders = [r[0] for r in conn.execute("SELECT processed_order FROM interactions ORDER BY processed_order")]
    assert orders == list(range(1, 13))


def test_failed_llm_does_not_touch_taxonomy(conn, fake):
    fake.fail_chat = True
    load_sample(conn, 3)
    before = conn.execute("SELECT COUNT(*) FROM taxonomy_nodes").fetchone()[0]
    _run(AnalysisWorker())
    run = conn.execute("SELECT * FROM analysis_runs ORDER BY id DESC").fetchone()
    assert run["status"] == "paused"
    assert conn.execute("SELECT COUNT(*) FROM taxonomy_nodes").fetchone()[0] == before
    assert conn.execute("SELECT COUNT(*) FROM interactions WHERE status='pending'").fetchone()[0] == 3


def test_proposals_require_review_when_auto_accept_disabled(conn, fake):
    db.save_settings(conn, {"batch_size": 30, "auto_accept_new_categories": False})
    load_sample(conn, 30)
    _run(AnalysisWorker())
    pending = taxonomy.list_changes(conn, status="proposed")
    creates = [p for p in pending if p["change_type"] == "create"]
    assert creates
    assert taxonomy.nodes(conn, ("active",), level=4) == []
    p = creates[0]
    with db.transaction(conn):
        res = taxonomy.accept(conn, p["id"], actor="user", settings=db.load_settings(conn))
    assert res["assigned"] >= 1
    assert taxonomy.get_node(conn, p["node_id"])["status"] == "active"
    if len(creates) > 1:
        with db.transaction(conn):
            taxonomy.reject(conn, creates[1]["id"])
        assert taxonomy.get_node(conn, creates[1]["node_id"])["status"] == "rejected"
        # rejected members are not re-proposed by the next re-clustering
        _run(AnalysisWorker(), job="recluster")
        again = [c for c in taxonomy.list_changes(conn, status="proposed") if c["change_type"] == "create"
                 and set(c["proposed"]["member_intent_ids"]) & set(creates[1]["proposed"]["member_intent_ids"])]
        assert not again


def test_pipeline_never_reads_ground_truth():
    for name in ("pipeline.py", "extraction.py", "taxonomy_builder.py", "clustering.py", "worker.py", "taxonomy.py",
                 "embeddings.py"):
        src = (APP_DIR / name).read_text()
        assert "ground_truth" not in src and "evaluation" not in src, name


def test_base_taxonomy_guard(conn):
    seed = json.loads((APP_DIR.parent / "data" / "taxonomy_seed.json").read_text())
    assert taxonomy.replace_base_taxonomy(conn, seed) > 1
    conn.execute("INSERT INTO interactions(conversation_id, seq, turns_json, created_at) VALUES ('Z', 1, '[]', 'x')")
    conn.execute("INSERT INTO intents(interaction_id, position, intent_text, created_at) VALUES (1, 0, 't', 'x')")
    with db.transaction(conn):
        assignments.create(conn, intent_id=1, interaction_id=1, l2_id="L2-BIL-INV", method="manual",
                           needs_review=False, review_reason=None)
    try:
        taxonomy.replace_base_taxonomy(conn, seed)
        assert False
    except taxonomy.TaxonomyError:
        pass


def test_semantically_distant_reuse_goes_to_review(conn, fake, monkeypatch):
    from app import clustering, embeddings, pipeline, taxonomy_builder

    db.save_settings(conn, {"batch_size": 10})
    load_sample(conn, 40)
    _run(AnalysisWorker())
    cents = embeddings.centroids(conn)
    l4s = [n for n in taxonomy.nodes(conn, ("active",), level=4) if n["id"] in cents]
    a, b = min(((x, y) for x in l4s for y in l4s if x["id"] < y["id"]),
               key=lambda p: float(cents[p[0]["id"]] @ cents[p[1]["id"]]))
    rows = conn.execute(
        "SELECT a.id AS aid, a.intent_id, i.intent_text, i.evidence, i.embedding_id, x.conversation_id "
        "FROM assignments a JOIN intents i ON i.id=a.intent_id JOIN interactions x ON x.id=a.interaction_id "
        "WHERE a.l4_id=?", (b["id"],)).fetchall()
    vecs = embeddings.load(conn, [r["embedding_id"] for r in rows])
    members = [{**dict(r), "vec": vecs[r["embedding_id"]]} for r in rows]

    asked = []

    def name_cluster(*args, exclude_l4=None, **kw):
        asked.append(exclude_l4)
        if not exclude_l4:
            return taxonomy_builder.Decision("existing", a["id"], "", "", None, "", "", "reuse")
        return taxonomy_builder.Decision("new", None, "Distinct Intent", "d", None, "Distinct Topic", "d", "new")

    monkeypatch.setattr(taxonomy_builder, "name_cluster", name_cluster)
    monkeypatch.setattr(clustering, "cluster_candidates", lambda X, t: clustering.ClusteringResult(
        [clustering.Cluster(list(range(len(X))), 0.9, "test")], [], "test"))
    l2 = pipeline._l2_lookup(conn)[assignments.path_for_node(conn, a["id"])["l2_id"]]
    pipeline.cluster_l2(pipeline.Ctx(conn, None, db.load_settings(conn), None), l2, members,
                        {"proposals": 0, "new_l4": 0, "new_l3": 0})

    assert len(asked) == 2 and a["id"] in asked[1]
    prop = conn.execute("SELECT * FROM taxonomy_changes WHERE change_type='create' AND status='proposed'").fetchone()
    assert prop and "first matched" in json.loads(prop["proposed_json"])["conflict"]
    assert not conn.execute("SELECT 1 FROM assignments WHERE l4_id=? AND intent_id IN (%s)" % ",".join(
        str(m["intent_id"]) for m in members), (a["id"],)).fetchone()
