"""Excel / CSV / JSON exports built from the current persisted state (including manual corrections)."""
from __future__ import annotations

import csv
import io
import json
import sqlite3
import zipfile
from collections import defaultdict

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from . import db, insights, interactions, taxonomy

SHEETS = ["Taxonomy", "Interaction Assignments", "Review Queue", "Taxonomy Changes"]

# Column headers, sheet names and fixed values per export language. Generated text (labels,
# descriptions, intents) stays in the analysis language.
TRANSLATIONS: dict[str, dict[str, str]] = {
    "de": {
        "Taxonomy": "Taxonomie", "Interaction Assignments": "Interaktionszuordnungen", "Review Queue": "Prüfliste",
        "Taxonomy Changes": "Taxonomie-Änderungen",
        "L1 ID": "L1-ID", "L2 ID": "L2-ID", "L3 ID": "L3-ID", "L4 ID": "L4-ID",
        "L1 Name": "L1-Name", "L2 Name": "L2-Name", "L3 Name": "L3-Name", "L4 Name": "L4-Name",
        "Level": "Ebene", "Source": "Quelle", "Description": "Beschreibung", "Interaction Count": "Anzahl Interaktionen",
        "Supporting Examples": "Belegende Beispiele", "Conversation ID": "Gesprächs-ID", "Channel": "Kanal",
        "Segment": "Segment", "Original Intent Summary": "Zusammenfassung", "Extracted Intent": "Extrahierter Intent",
        "Evidence Excerpt": "Belegstelle", "Similarity": "Ähnlichkeit", "Classification Method": "Zuordnungsmethode",
        "Review Status": "Prüfstatus", "Reviewer Notes": "Prüfnotizen", "Queue Reason": "Grund",
        "Current Path": "Aktueller Pfad", "Suggested Correction": "Vorgeschlagene Korrektur",
        "Change Timestamp": "Zeitpunkt", "Category ID": "Kategorie-ID", "Change Type": "Änderungsart",
        "Previous Value": "Vorheriger Wert", "Proposed / New Value": "Vorgeschlagener / neuer Wert",
        "Approval Status": "Freigabestatus", "Rationale": "Begründung", "Decided At": "Entschieden am",
        "Taxonomy Version": "Taxonomie-Version", "(empty)": "(leer)",
        # values
        "pending": "Ungeprüft", "needs_review": "Zu prüfen", "confirmed": "Bestätigt", "corrected": "Korrigiert",
        "unclassified": "Nicht klassifiziert", "failed": "Fehlgeschlagen", "processed": "Verarbeitet",
        "predefined": "fest", "ai": "KI-generiert", "user": "manuell",
        "create": "Angelegt", "update": "Bearbeitet", "refine": "Beschreibung verfeinert", "merge": "Zusammenführung",
        "split": "Aufteilung", "retire": "Stillgelegt", "move": "Verschoben",
        "proposed": "vorgeschlagen", "accepted": "übernommen", "auto_accepted": "automatisch übernommen",
        "rejected": "abgelehnt",
        "llm_extraction": "LLM-Extraktion", "embedding_match": "Embedding-Treffer", "cluster": "Cluster",
        "llm_cluster_match": "LLM-Clusterzuordnung", "llm_outlier": "LLM (seltener Intent)", "manual": "Manuell",
        "candidate": "Kandidat",
        "Processing failed": "Verarbeitung fehlgeschlagen", "No customer intent extracted": "Kein Kundenanliegen erkannt",
        "Unclassified": "Nicht klassifiziert", "No matching L1/L2 category": "Keine passende L1/L2-Kategorie",
        "No L3/L4 category yet": "Noch keine L3/L4-Kategorie", "Low similarity": "Geringe Ähnlichkeit",
        "Flagged by reviewer": "Von Prüfer markiert", "Awaiting validation": "Wartet auf Bestätigung",
    },
}
VALUE_COLUMNS = {"Source", "Review Status", "Classification Method", "Change Type", "Approval Status"}


def tr(text: str, lang: str) -> str:
    return TRANSLATIONS.get(lang, {}).get(text, text)


def _localize(rows: list[dict], lang: str) -> list[dict]:
    if lang not in TRANSLATIONS:
        return rows
    out = []
    for r in rows:
        row = {}
        for k, v in r.items():
            if k in VALUE_COLUMNS and isinstance(v, str):
                v = tr(v, lang)
            elif k == "Queue Reason" and isinstance(v, str):
                v = "; ".join(_reason(part, lang) for part in v.split("; "))
            row[tr(k, lang)] = v
        out.append(row)
    return out


def _reason(part: str, lang: str) -> str:
    if part.startswith("Low similarity"):
        return tr("Low similarity", lang) + part[len("Low similarity"):]
    return tr(part, lang)


def _node_stats(conn: sqlite3.Connection) -> tuple[dict, dict]:
    counts: dict[str, set] = defaultdict(set)
    examples: dict[str, list[str]] = defaultdict(list)
    for r in conn.execute("SELECT a.interaction_id, a.l1_id, a.l2_id, a.l3_id, a.l4_id, i.intent_text, i.evidence "
                          "FROM assignments a JOIN intents i ON i.id=a.intent_id WHERE a.review_status!='unclassified' "
                          "ORDER BY a.similarity DESC"):
        for k in ("l1_id", "l2_id", "l3_id", "l4_id"):
            if r[k]:
                counts[r[k]].add(r["interaction_id"])
                if len(examples[r[k]]) < 3:
                    examples[r[k]].append(r["evidence"] or r["intent_text"])
    return counts, examples


def taxonomy_rows(conn: sqlite3.Connection) -> list[dict]:
    nodes = {n["id"]: n for n in taxonomy.nodes(conn, ("active",))}
    children: dict[str | None, list[dict]] = defaultdict(list)
    for n in sorted(nodes.values(), key=lambda n: n["label"]):
        children[n["parent_id"]].append(n)
    counts, examples = _node_stats(conn)
    rows = []

    def emit(path: list[dict]):
        leaf = path[-1]
        row = {}
        for lvl in range(1, 5):
            n = path[lvl - 1] if len(path) >= lvl else None
            row[f"L{lvl} ID"] = n["id"] if n else ""
            row[f"L{lvl} Name"] = n["label"] if n else ""
        row["Level"] = leaf["level"]
        row["Source"] = leaf["source"]
        row["Description"] = leaf["description"]
        row["Interaction Count"] = len(counts.get(leaf["id"], ()))
        row["Supporting Examples"] = " | ".join(examples.get(leaf["id"], []))
        rows.append(row)

    def walk(node: dict, path: list[dict]):
        path = path + [node]
        kids = children.get(node["id"], [])
        if not kids or node["level"] == 4:
            emit(path)
            return
        for k in kids:
            walk(k, path)

    for root in children.get(None, []):
        walk(root, [])
    return rows


def assignment_rows(conn: sqlite3.Connection) -> list[dict]:
    rows = []
    for inter in interactions.list_interactions(conn):
        notes = inter.get("reviewer_notes") or ""
        full = {a["id"]: a for a in _assignments_full(conn, inter["id"])}
        base = {"Conversation ID": inter["conversation_id"], "Channel": inter["channel"] or "",
                "Segment": inter["segment"] or "", "Original Intent Summary": inter["summary"] or ""}
        if not inter["assignments"]:
            rows.append({**base, "Extracted Intent": "", "L1": "", "L2": "", "L3": "", "L4": "",
                         "L1 ID": "", "L2 ID": "", "L3 ID": "", "L4 ID": "", "Evidence Excerpt": "",
                         "Similarity": "", "Classification Method": "",
                         "Review Status": inter["review_status"] if inter["status"] == "processed" else inter["status"],
                         "Reviewer Notes": notes})
            continue
        for a in inter["assignments"]:
            f = full[a["id"]]
            rows.append({**base, "Extracted Intent": a["intent_text"],
                         "L1": a["l1_label"] or "", "L2": a["l2_label"] or "", "L3": a["l3_label"] or "",
                         "L4": a["l4_label"] or "", "L1 ID": a["l1_id"] or "", "L2 ID": a["l2_id"] or "",
                         "L3 ID": a["l3_id"] or "", "L4 ID": a["l4_id"] or "", "Evidence Excerpt": f["evidence"] or "",
                         "Similarity": round(a["similarity"], 3) if a["similarity"] is not None else "",
                         "Classification Method": a["method"], "Review Status": a["review_status"],
                         "Reviewer Notes": " / ".join(x for x in (f["reviewer_notes"], notes) if x)})
    return rows


def _assignments_full(conn: sqlite3.Connection, interaction_id: int) -> list[sqlite3.Row]:
    return conn.execute("SELECT a.id, a.reviewer_notes, a.review_reason, a.alternatives_json, i.evidence FROM assignments a "
                        "JOIN intents i ON i.id=a.intent_id WHERE a.interaction_id=?", (interaction_id,)).fetchall()


def review_rows(conn: sqlite3.Connection) -> list[dict]:
    settings = db.load_settings(conn)
    match = settings.thresholds()["match"]
    labels = {r["id"]: r["label"] for r in conn.execute("SELECT id, label FROM taxonomy_nodes")}
    rows = []
    for inter in interactions.list_interactions(conn):
        full = {a["id"]: a for a in _assignments_full(conn, inter["id"])}
        if inter["status"] == "failed":
            rows.append({"Conversation ID": inter["conversation_id"], "Queue Reason": "Processing failed",
                         "Extracted Intent": "", "Current Path": "", "Similarity": "", "Suggested Correction": "",
                         "Review Status": "failed"})
        if inter["status"] == "processed" and not inter["assignments"]:
            rows.append({"Conversation ID": inter["conversation_id"], "Queue Reason": "No customer intent extracted",
                         "Extracted Intent": "", "Current Path": "", "Similarity": "", "Suggested Correction": "",
                         "Review Status": inter["review_status"]})
        for a in inter["assignments"]:
            if a["review_status"] in ("confirmed", "corrected"):
                continue
            reasons = []
            f = full[a["id"]]
            if a["review_status"] == "unclassified":
                reasons.append("Unclassified")
            if not a["l2_id"]:
                reasons.append("No matching L1/L2 category")
            elif not a["l4_id"]:
                reasons.append("No L3/L4 category yet")
            if a["needs_review"] and f["review_reason"]:
                reasons.append(f["review_reason"])
            if a["similarity"] is not None and a["l4_id"] and a["similarity"] < match:
                reasons.append(f"Low similarity ({a['similarity']:.2f})")
            if inter["review_flag"]:
                reasons.append("Flagged by reviewer")
            if not reasons:
                reasons.append("Awaiting validation")
            alts = db.loads(f["alternatives_json"], [])
            suggestion = ", ".join(f"{labels.get(x['node_id'], x['node_id'])} ({x['similarity']:.2f})"
                                   for x in alts[:2] if x.get("node_id") != a["l4_id"])
            path = " > ".join(x for x in (a["l1_label"], a["l2_label"], a["l3_label"], a["l4_label"]) if x)
            rows.append({"Conversation ID": inter["conversation_id"], "Queue Reason": "; ".join(reasons),
                         "Extracted Intent": a["intent_text"], "Current Path": path,
                         "Similarity": round(a["similarity"], 3) if a["similarity"] is not None else "",
                         "Suggested Correction": suggestion, "Review Status": a["review_status"]})
    rows.sort(key=lambda r: r["Queue Reason"] == "Awaiting validation")
    return rows


def change_rows(conn: sqlite3.Connection) -> list[dict]:
    rows = []
    for c in reversed(taxonomy.list_changes(conn, limit=100000)):
        rows.append({"Change Timestamp": c["created_at"], "Category ID": c["node_id"] or "",
                     "Change Type": c["change_type"],
                     "Previous Value": json.dumps(c["previous"], ensure_ascii=False) if c["previous"] else "",
                     "Proposed / New Value": json.dumps(_compact(c["proposed"]), ensure_ascii=False),
                     "Approval Status": c["status"], "Source": c["source"], "Rationale": c["rationale"] or "",
                     "Decided At": c["decided_at"] or "", "Taxonomy Version": c["version"] or ""})
    return rows


def _compact(p: dict) -> dict:
    return {k: v for k, v in p.items() if k not in ("member_intent_ids",)} | (
        {"members": len(p["member_intent_ids"])} if "member_intent_ids" in p else {})


def all_sheets(conn: sqlite3.Connection, lang: str = "en") -> dict[str, list[dict]]:
    """Sheets keyed by their English id; rows use headers and fixed values in ``lang``."""
    return {"Taxonomy": _localize(taxonomy_rows(conn), lang),
            "Interaction Assignments": _localize(assignment_rows(conn), lang),
            "Review Queue": _localize(review_rows(conn), lang),
            "Taxonomy Changes": _localize(change_rows(conn), lang)}


def _file_name(sheet: str, lang: str) -> str:
    name = tr(sheet, lang).lower().translate(str.maketrans({"ä": "ae", "ö": "oe", "ü": "ue", "ß": "ss", " ": "_"}))
    return name + ".csv"


def to_xlsx(conn: sqlite3.Connection, lang: str = "en") -> bytes:
    wb = Workbook()
    wb.remove(wb.active)
    header_fill = PatternFill("solid", fgColor="F1F5F9")
    for name, rows in all_sheets(conn, lang).items():
        ws = wb.create_sheet(tr(name, lang)[:31])
        cols = list(rows[0].keys()) if rows else [tr("(empty)", lang)]
        ws.append(cols)
        for c in ws[1]:
            c.font = Font(bold=True)
            c.fill = header_fill
        for r in rows:
            ws.append([r.get(c, "") for c in cols])
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        for i, col in enumerate(cols, 1):
            width = max([len(str(col))] + [min(60, len(str(r.get(col, "")))) for r in rows[:300]]) + 2
            ws.column_dimensions[get_column_letter(i)].width = min(62, max(10, width))
        for row in ws.iter_rows(min_row=2):
            for c in row:
                c.alignment = Alignment(vertical="top", wrap_text=isinstance(c.value, str) and len(c.value) > 60)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _csv(rows: list[dict]) -> str:
    buf = io.StringIO()
    if rows:
        w = csv.DictWriter(buf, fieldnames=list(rows[0].keys()))
        w.writeheader()
        w.writerows(rows)
    return buf.getvalue()


def to_csv(conn: sqlite3.Connection, sheet: str | None, lang: str = "en") -> tuple[bytes, str, str]:
    if sheet and sheet not in SHEETS:
        raise KeyError(sheet)
    sheets = all_sheets(conn, lang)
    if sheet:
        return ("\ufeff" + _csv(sheets[sheet])).encode("utf-8"), "text/csv", _file_name(sheet, lang)
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as z:
        for name, rows in sheets.items():
            z.writestr(_file_name(name, lang), "\ufeff" + _csv(rows))
    return buf.getvalue(), "application/zip", "intent_studio_export_csv.zip"


def to_json(conn: sqlite3.Connection) -> dict:
    return {
        "exported_at": db.utcnow(),
        "summary": insights.summary(conn),
        "taxonomy": taxonomy.tree(conn, include_inactive=True)["nodes"],
        "taxonomy_rows": taxonomy_rows(conn),
        "assignments": assignment_rows(conn),
        "review_queue": review_rows(conn),
        "taxonomy_changes": change_rows(conn),
    }
